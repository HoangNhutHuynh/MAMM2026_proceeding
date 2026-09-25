#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rebuild program.json for the MAMM 2026 programme website.

Two sources, either or both:

  * the technical-program PDF  -> schedule, sessions, chairs, keynote and
    invited speakers, and which paper is in which session
  * the Word proceedings file  -> paper titles, authors, affiliations,
    abstracts, keywords, committees, preface

    pip install python-docx pdfplumber
    python3 tools/build_program.py source/Technical.pdf "source/Proceedings MAMM_revise.docx"

Give the files in any order; they are told apart by extension. With only the
.docx, the schedule is read from the Word tables instead (the original
behaviour). With only the .pdf, papers have no abstracts.

Papers listed in the PDF but missing from the Word file are kept as
placeholders ("Title to be announced") so they still appear in the schedule.
Fill them in without waiting for a new Word file by creating
source/extra-papers.json:

    {"papers": {"ID_91": {"title": "...", "abstract": "...",
                          "keywords": ["..."],
                          "authors": [{"name": "...", "affiliations": ["..."]}]}}}

Options:
    -o PATH        output file (default: program.json beside index.html)
    --keep-hyphens do not rejoin words broken by justified hyphenation
    --check-only   validate and report, write nothing
"""

import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pdf_program

try:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")


# --------------------------------------------------------------------- config
DATEMAP = {'5 Nov': '2026-11-05', '6 Nov': '2026-11-06', '7 Nov': '2026-11-07'}
DAYMAP = {'5 Nov': 1, '6 Nov': 2, '7 Nov': 3}
DAY_DATE = {1: '2026-11-05', 2: '2026-11-06', 3: '2026-11-07'}

# Poster sessions have no time/room printed in their own table header.
POSTER_TIME = {1: '09:20-10:50 / 14:40-16:10', 2: '09:20-10:50 / 14:40-16:10'}
POSTER_ROOM = {1: 'Room 1', 2: 'Rooms 1, 2 & 3'}

# Genuine hyphenated compounds that de-hyphenation must not merge.
# Add to this list if the report below flags a word that should stay hyphenated.
KEEP_HYPHENATED = {'counter-ions', 'multi-mode', 'non-invasive', 'pre-clinical'}

CONFERENCE = {
    'name': 'MAMM 2026',
    'fullName': 'The 8th International Conference on Microactuators, '
                'Microsensors and Micromechanisms',
    'venue': "Sun Moon Lake Teachers' Hostel, Nantou, Taiwan",
    'dates': '5-7 November 2026',
    'startDate': '2026-11-05',
    'endDate': '2026-11-07',
    'host': 'National Chung Hsing University, Taiwan',
}

STOP_HEADINGS = {'ORGANIZED BY', 'KEYNOTE SPEAKER', 'INVITED SPEAKER', 'TECHNICAL PROGRAM'}

warnings = []


def warn(msg):
    warnings.append(msg)


# ----------------------------------------------------------------- doc reading
def read_body(path):
    """Return the document body as an ordered list of ('p', style, text) and
    ('t', None, rows) items, where rows is a list of lists of cell paragraphs."""
    d = docx.Document(path)
    items = []
    for child in d.element.body.iterchildren():
        if child.tag == qn('w:p'):
            p = Paragraph(child, d)
            t = p.text.strip()
            if t:
                items.append(('p', p.style.name, t))
        elif child.tag == qn('w:tbl'):
            tb = Table(child, d)
            rows = []
            for r in tb.rows:
                cells, seen = [], set()
                for c in r.cells:                       # skip merged duplicates
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append([x.text.strip() for x in c.paragraphs if x.text.strip()])
                rows.append(cells)
            items.append(('t', None, rows))
    return items


def flat(cell):
    return ' / '.join(cell)


def first_cell(rows):
    return flat(rows[0][0]) if rows and rows[0] else ''


# ------------------------------------------------------------- table locating
RE_DAY_TABLE = re.compile(r'MAMM\s*\d{4}\s*-\s*DAY\s*(\d)', re.I)
RE_SESSION_HEAD = re.compile(r'^(Oral|Poster) Session \((\d+)\)')


def classify_tables(items):
    days, detail, summary, organizers = [], [], [], None
    for kind, _, rows in items:
        if kind != 't':
            continue
        head = first_cell(rows)
        m = RE_DAY_TABLE.search(head)
        if m:
            days.append((int(m.group(1)), rows))
            continue
        if RE_SESSION_HEAD.match(head):
            ncols = max(len(r) for r in rows)
            (detail if ncols >= 3 else summary).append(rows)
            continue
        if organizers is None and 'Organizers' in head:
            organizers = rows
    return days, detail, summary, organizers


# --------------------------------------------------------------- day schedule
ROOM_RE = re.compile(r'Room\s*([0-9])')


def room_of(text, idx, ncols):
    if re.search(r'Rooms\s*1\s*,?\s*2\s*&\s*3', text):
        return 'All rooms'
    if 'Outside Room' in text:
        return 'Outside Room 1'
    rooms = ROOM_RE.findall(text)
    if rooms:
        return 'Room ' + rooms[0]
    if ncols == 1:
        return 'All'
    return 'Room %d' % (idx + 1)


def parse_days(day_tables):
    days = []
    for dayno, rows in sorted(day_tables):
        title = first_cell(rows)
        slots = []
        for r in rows[1:]:
            if not r:
                continue
            time = flat(r[0])
            cells = [flat(c) for c in r[1:] if flat(c)]
            n = len(cells)
            for i, c in enumerate(cells):
                slots.append({'time': time, 'room': room_of(c, i, n),
                              'text': c, 'col': i, 'ncols': n})
        days.append({'day': dayno, 'date': DAY_DATE.get(dayno), 'title': title, 'slots': slots})
    return days


# ------------------------------------------------------------------- sessions
def parse_session_header(paras, kind):
    m = RE_SESSION_HEAD.match(paras[0])
    num = int(m.group(2))
    rest = paras[0][m.end():].lstrip(' -\u2013').strip()
    chairs = date = time = room = None
    name = rest
    for p in paras[1:]:
        if p.lower().startswith('chairs'):
            chairs = p.split(':', 1)[1].strip()
        elif re.search(r'\d{1,2}:\d{2}', p):
            mm = re.match(r'(\d+ \w+),\s*([\d:\-\u2013, /]+)\s*\((.*)\)', p)
            if mm:
                date, time, room = mm.group(1), mm.group(2).strip(), mm.group(3).strip()
    if kind == 'poster':
        mm = re.search(r'Day (\d).*\((\d+ \w+)\)', rest)
        if mm:
            name = 'Poster Session %d' % num
            date = mm.group(2)
        time, room = POSTER_TIME.get(num, time), POSTER_ROOM.get(num, room)
    return {'num': num, 'name': name, 'chairs': chairs,
            'date': date, 'time': time, 'room': room}


def parse_sessions(detail_tables):
    sessions, sched = [], {}
    for rows in detail_tables:
        head_paras = rows[0][0]
        kind = 'oral' if head_paras[0].startswith('Oral') else 'poster'
        h = parse_session_header(head_paras, kind)
        sid = ('OS' if kind == 'oral' else 'PS') + str(h['num'])
        if h['time'] is None or h['room'] is None:
            warn('%s: could not read time/room from its table header' % sid)
        ids = []
        for r in rows[1:]:
            pid = flat(r[0]).strip()
            if not pid.startswith('ID_'):
                continue                                 # header row 'ID | Title | Author(s)'
            prov = '*' in pid
            pid = pid.replace('*', '').strip()
            if pid in sched:
                warn('%s appears in more than one session (%s and %s)'
                     % (pid, sched[pid]['session'], sid))
            ids.append(pid)
            sched[pid] = {'session': sid, 'provisional': prov}
        if not ids:
            warn('%s: no papers found in its table' % sid)
        sessions.append({
            'id': sid, 'kind': kind, 'num': h['num'], 'name': h['name'],
            'chairs': h['chairs'],
            'date': DATEMAP.get(h['date']), 'dateLabel': h['date'],
            'day': DAYMAP.get(h['date']), 'time': h['time'], 'room': h['room'],
            'paperIds': ids,
        })
    sessions.sort(key=lambda s: (s['kind'] != 'oral', s['num']))
    return sessions, sched


# ------------------------------------------------------------------ abstracts
def parse_abstracts(items):
    papers, cur, mode = {}, None, None
    for kind, style, val in items:
        if kind != 'p':
            continue
        if style == 'Heading 1':
            m = re.match(r'(ID_\d+):\s*(.*)', val)
            if m:
                cur = {'id': m.group(1), 'title': m.group(2).strip(), 'authors': [],
                       'affiliations': {}, 'abstract': '', 'keywords': [], 'authorsLine': ''}
                if cur['id'] in papers:
                    warn('%s: duplicate abstract entry' % cur['id'])
                papers[cur['id']] = cur
                mode = 'authors'
            else:
                cur, mode = None, None
            continue
        if cur is None:
            continue
        if style == 'Heading 2' and mode == 'authors':
            cur['authorsLine'] = val
            for part in re.split(r',| and ', val):
                part = part.strip()
                if not part:
                    continue
                mm = re.match(r'^(.*?)[\s]*([\d\*\u2020,]*)$', part)
                nm = mm.group(1).strip().rstrip('*').strip()
                refs = re.findall(r'\d+', mm.group(2) or '')
                if nm:
                    cur['authors'].append({'name': nm, 'refs': refs})
        elif style == 'Affiliation':
            mm = re.match(r'^(\d+)\s+(.*)$', val)
            if mm:
                cur['affiliations'][mm.group(1)] = mm.group(2).strip()
            else:
                cur['affiliations'].setdefault('1', val)
        elif val.strip().lower() == 'abstract':
            mode = 'abstract'
        elif val.lower().startswith('keywords'):
            kw = val.split(':', 1)[1] if ':' in val else ''
            cur['keywords'] = [k.strip().rstrip('.') for k in re.split(r'[,;]', kw) if k.strip()]
            mode = None
        elif mode == 'abstract':
            cur['abstract'] = (cur['abstract'] + ' ' + val).strip()
    return papers


# --------------------------------------------------------------- text cleanup
EMAIL = re.compile(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}')
PHONE = re.compile(r'\b0\d{8,10}\b')


def scrub(txt):
    """Strip e-mail addresses and phone numbers that appear as affiliations."""
    txt = PHONE.sub('', EMAIL.sub('', txt))
    txt = re.sub(r'\(\s*[;,\s]*\)', '', txt)
    txt = re.sub(r'^[\*\s]*E-?mail\s*:?\s*[;,\s]*$', '', txt, flags=re.I)
    return re.sub(r'\s{2,}', ' ', txt).strip(' ;,')


TOK = re.compile(r'[A-Za-z]+(?:-[A-Za-z]+)*')
HYP = re.compile(r'\b([A-Za-z]{2,})-([a-z]{2,})\b')


def dehyphenate(papers):
    """Rejoin words Word broke across lines ('sys-tem' -> 'system').

    A hyphenated pair is merged only when the joined form occurs unhyphenated
    somewhere in the corpus AND at least one half never occurs as a word on its
    own -- so real compounds like 'bridge-type' are left alone."""
    vocab = set()
    for p in papers:
        for w in TOK.findall(' '.join([p['abstract'], p['title']] + p['keywords'])):
            if '-' not in w:
                vocab.add(w.lower())

    merged = set()

    def rep(m):
        if m.group(0).lower() in KEEP_HYPHENATED:
            return m.group(0)
        a, b = m.group(1), m.group(2)
        if (a + b).lower() in vocab and (a.lower() not in vocab or b.lower() not in vocab):
            merged.add(m.group(0))
            return a + b
        return m.group(0)

    for p in papers:
        p['abstract'] = HYP.sub(rep, p['abstract'])
    return sorted(merged)


# ----------------------------------------------------------------- committees
def parse_committees(items):
    committees, cur = [], None
    for kind, _, val in items[:120]:
        if kind != 'p':
            continue
        if val.rstrip(':') in STOP_HEADINGS:
            break
        if val.isupper() and not val.startswith('_') and len(val) < 60 and 'PREFACE' not in val:
            cur = {'name': val.rstrip(':'), 'members': []}
            committees.append(cur)
        elif val.startswith('_') and cur:
            nm, _, af = val.lstrip('_').strip().partition(',')
            cur['members'].append({'name': nm.strip(), 'affiliation': af.strip()})
    return committees


# ----------------------------------------------------------------- validation
def validate_docx_lists(sessions, summary_tables):
    """Cross-check each session against the document's own summary tables."""
    ok = True
    expected = {}
    for rows in summary_tables:
        for r in rows:
            if len(r) < 2:
                continue
            m = RE_SESSION_HEAD.match(flat(r[0]))
            if not m:
                continue
            sid = ('OS' if m.group(1) == 'Oral' else 'PS') + m.group(2)
            expected[sid] = set(x.strip() for x in flat(r[1]).split(',') if x.strip())
    if not expected:
        return None
    for s in sessions:
        exp = expected.get(s['id'])
        if exp is None:
            warn('%s: no matching row in the ORAL/POSTER Presentation List tables' % s['id'])
            ok = False
            continue
        got = set(s['paperIds'])
        if exp != got:
            ok = False
            if exp - got:
                warn('%s: in the summary list but missing from the detail table: %s'
                     % (s['id'], ', '.join(sorted(exp - got))))
            if got - exp:
                warn('%s: in the detail table but missing from the summary list: %s'
                     % (s['id'], ', '.join(sorted(got - exp))))
    return ok


def validate_common(sessions, papers_out):
    seen = {}
    for s in sessions:
        if not s['paperIds']:
            warn('%s: no papers assigned' % s['id'])
        if not s.get('time') or not s.get('room'):
            warn('%s: time or room could not be read' % s['id'])
        for pid in s['paperIds']:
            if pid in seen:
                warn('%s is listed in two sessions: %s and %s' % (pid, seen[pid], s['id']))
            seen[pid] = s['id']
    for p in papers_out:
        if p.get('pending'):
            continue
        if len(p['abstract']) < 50:
            warn('%s: abstract missing or very short' % p['id'])
        if not p['keywords']:
            warn('%s: no keywords' % p['id'])
        if not p['authors']:
            warn('%s: no authors' % p['id'])


# ------------------------------------------------------------------ main flow
def schedule_from_docx(items):
    day_tables, detail_tables, summary_tables, org_rows = classify_tables(items)
    if not day_tables:
        sys.exit('No day-schedule tables found in the Word file, and no PDF was given. '
                 'Expected a table whose first cell reads like '
                 '"MAMM 2026 - DAY 1 - 5 NOVEMBER 2026 (THURSDAY)".')
    if not detail_tables:
        sys.exit('No session tables found in the Word file. Expected tables whose first cell '
                 'starts with "Oral Session (1) - ..." or "Poster Session (1) - ...".')
    days = parse_days(day_tables)
    sessions, sched = parse_sessions(detail_tables)
    return days, sessions, sched, summary_tables, org_rows


def load_extras(pdf_path, docx_path, explicit=None):
    """source/extra-papers.json, if present, fills in papers the Word file lacks."""
    cands = []
    if explicit:
        cands.append(explicit)
    for f in (pdf_path, docx_path):
        if f:
            cands.append(os.path.join(os.path.dirname(os.path.abspath(f)), 'extra-papers.json'))
    here = os.path.dirname(os.path.abspath(__file__))
    cands.append(os.path.join(os.path.dirname(here), 'source', 'extra-papers.json'))
    for c in cands:
        if c and os.path.exists(c):
            with open(c, encoding='utf-8') as fh:
                data = json.load(fh)
            return data.get('papers', data), c
    return {}, None


def build(pdf_path=None, docx_path=None, keep_hyphens=False, extras_path=None):
    items = read_body(docx_path) if docx_path else []
    abstracts = parse_abstracts(items) if items else {}
    extras, extras_file = load_extras(pdf_path, docx_path, extras_path)

    summary_tables, org_rows, lists_ok = [], None, None
    if pdf_path:
        sched_src = 'PDF'
        pdfdata = pdf_program.parse(pdf_path, warn)
        days, sessions = pdfdata['days'], pdfdata['sessions']
        provisional = {}
        if items:
            _, _, summary_tables, org_rows = classify_tables(items)
    else:
        sched_src = 'Word'
        days, sessions, sched, summary_tables, org_rows = schedule_from_docx(items)
        provisional = dict((k, v['provisional']) for k, v in sched.items())
        lists_ok = validate_docx_lists(sessions, summary_tables)

    where = {}
    for s in sessions:
        for pid in s['paperIds']:
            where.setdefault(pid, s)

    out_papers, pending = [], []
    for pid, s in sorted(where.items(), key=lambda kv: int(kv[0].split('_')[1])):
        a = abstracts.get(pid)
        x = extras.get(pid) or {}
        if a is None and not x:
            pending.append(pid)
        title = x.get('title') or (a['title'] if a else 'Title to be announced')
        if a:
            authors = []
            for au in a['authors']:
                au['affiliations'] = [a['affiliations'][r] for r in au['refs']
                                      if r in a['affiliations']]
                authors.append({'name': au['name'], 'affiliations': au['affiliations']})
        else:
            authors = []
        if x.get('authors'):
            authors = [{'name': au.get('name', ''),
                        'affiliations': au.get('affiliations', [])} for au in x['authors']]
        affs = []
        for au in authors:
            for y in au['affiliations']:
                if y not in affs:
                    affs.append(y)
        if a and not affs:
            affs = list(a['affiliations'].values())
        out_papers.append({
            'id': pid, 'num': int(pid.split('_')[1]),
            'title': title,
            'authors': authors,
            'authorsLine': (a or {}).get('authorsLine', ''),
            'affiliations': affs,
            'abstract': x.get('abstract') or (a['abstract'] if a else ''),
            'keywords': x.get('keywords') or (a['keywords'] if a else []),
            'type': s['kind'],
            'sessionId': s['id'], 'sessionName': s['name'], 'sessionNum': s['num'],
            'day': s['day'], 'date': s['date'], 'time': s['time'], 'room': s['room'],
            'chairs': s.get('chairs'),
            'provisional': bool(provisional.get(pid)),
            'pending': a is None and not x,
        })

    for pid in abstracts:
        if pid not in where:
            warn('%s has an abstract in the Word file but is not in the schedule' % pid)

    merged = [] if keep_hyphens else dehyphenate([p for p in out_papers if p['abstract']])

    for p in out_papers:
        p['affiliations'] = [y for y in (scrub(z) for z in p['affiliations']) if len(y) > 3]
        for au in p['authors']:
            au['affiliations'] = [y for y in (scrub(z) for z in au['affiliations']) if len(y) > 3]

    validate_common(sessions, out_papers)

    organizers = [flat(c) for r in (org_rows or []) for c in r if flat(c)]
    preface = [v for k, st, v in items[:12] if k == 'p' and st == 'Normal' and len(v) > 150]

    conf = dict(CONFERENCE)
    conf['organizers'] = organizers
    conf['preface'] = preface

    notes = []
    if any(p['provisional'] for p in out_papers):
        notes.append('Papers marked with an asterisk (*) in the professor-selection groups are a '
                     'provisional pick pending confirmation from the corresponding professor.')
    if pending:
        notes.append('Details for %s will be announced.' % ', '.join(pending))
    notes.append('Keynote and invited speaker slots follow the numbering printed in the '
                 'official technical programme.')

    data = {
        'conference': conf,
        'days': days,
        'sessions': sessions,
        'papers': out_papers,
        'committees': parse_committees(items) if items else [],
        'notes': notes,
    }
    info = {'source': sched_src, 'lists_ok': lists_ok, 'merged': merged,
            'pending': pending, 'extras_file': extras_file}
    return data, info


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    default_out = os.path.join(os.path.dirname(here), 'program.json')

    ap = argparse.ArgumentParser(description='Rebuild program.json from the programme PDF '
                                             'and/or the Word proceedings file.')
    ap.add_argument('files', nargs='+', help='the .pdf programme and/or the .docx proceedings')
    ap.add_argument('-o', '--out', default=default_out, help='output JSON path')
    ap.add_argument('--extras', help='path to extra-papers.json')
    ap.add_argument('--keep-hyphens', action='store_true',
                    help='do not rejoin words broken by justified hyphenation')
    ap.add_argument('--check-only', action='store_true', help='validate only, write nothing')
    args = ap.parse_args()

    pdf_path = docx_path = None
    for f in args.files:
        ext = os.path.splitext(f)[1].lower()
        if ext == '.pdf':
            pdf_path = f
        elif ext in ('.docx', '.dotx'):
            docx_path = f
        else:
            sys.exit('Do not know what to do with %s (expected .pdf or .docx)' % f)
    if not pdf_path and not docx_path:
        sys.exit('Give at least one .pdf or .docx file.')

    data, info = build(pdf_path, docx_path, args.keep_hyphens, args.extras)

    oral = [p for p in data['papers'] if p['type'] == 'oral']
    poster = [p for p in data['papers'] if p['type'] == 'poster']
    nums = sorted(p['num'] for p in data['papers'])
    gaps = [n for n in range(1, (nums[-1] if nums else 0) + 1) if n not in set(nums)]

    print('-' * 64)
    print('Schedule from : %s' % info['source'])
    print('Abstracts from: %s' % (os.path.basename(docx_path) if docx_path else 'none'))
    if info['extras_file']:
        print('Extra papers  : %s' % info['extras_file'])
    print('Papers        : %d  (%d oral, %d poster)' % (len(data['papers']), len(oral), len(poster)))
    print('Sessions      : %d oral, %d poster'
          % (sum(1 for s in data['sessions'] if s['kind'] == 'oral'),
             sum(1 for s in data['sessions'] if s['kind'] == 'poster')))
    print('Days          : %d' % len(data['days']))
    print('Committees    : %d' % len(data['committees']))
    prov = [p['id'] for p in data['papers'] if p['provisional']]
    if prov:
        print('Provisional   : %s' % ', '.join(prov))
    if gaps:
        print('Unused IDs    : %s' % ', '.join('ID_%02d' % n for n in gaps))
    if info['pending']:
        print('AWAITING TITLE/ABSTRACT: %s' % ', '.join(info['pending']))
        print('              (add them to source/extra-papers.json, or to the Word file)')
    if info['merged']:
        print('De-hyphenated %d word(s): %s' % (len(info['merged']), ', '.join(info['merged'][:10])
              + (' ...' if len(info['merged']) > 10 else '')))
    if info['lists_ok'] is not None:
        print('Session lists match the summary tables: %s' % ('YES' if info['lists_ok'] else 'NO'))

    if warnings:
        print('\n%d warning(s):' % len(warnings))
        for w in warnings:
            print('  ! ' + w)
    else:
        print('\nNo warnings.')
    print('-' * 64)

    if args.check_only:
        print('--check-only: nothing written.')
        return 0 if not warnings else 1

    with open(args.out, 'w', encoding='utf-8') as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)
    print('Wrote %s (%.0f KB)' % (args.out, os.path.getsize(args.out) / 1024.0))
    print('Now commit program.json and push -- GitHub Pages redeploys automatically.')
    return 0


if __name__ == '__main__':
    sys.exit(main())

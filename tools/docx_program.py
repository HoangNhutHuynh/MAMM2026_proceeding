# -*- coding: utf-8 -*-
"""
Read the technical programme that is embedded in the Word proceedings file
("Proceedings MAMM_revise.docx") and return the same structure that
pdf_program.parse() returns for the PDF:

    {'days': [...], 'sessions': [...], 'speakers': [...]}

What it reads:

  * the day tables         first cell "MAMM 2026 - DAY 1 - 5 NOVEMBER 2026 (THURSDAY)";
                           several days may share one Word table
  * the session tables     first cell "Oral Session (1) - <name>", then
                           "Chairs: A & B" and "5 Nov, 09:20-10:50 (Room 1)",
                           followed by one row per paper (ID | Title | Authors)
  * the speaker tables     one-column tables "Prof. X / affiliation",
                           "Topic: ...", "Abstract: ..." under a KEYNOTE SPEAKER or
                           INVITED SPEAKER heading

Merged cells are read from the Word XML directly (gridSpan / vMerge), so a
cell that spans two time rows is reported once, with the combined time range.

    python3 tools/docx_program.py "source/Proceedings MAMM_revise.docx"

prints what was found -- the quickest way to see why a parse went wrong.
"""
import re
import sys

try:
    import docx
    from docx.oxml.ns import qn
except ImportError:
    docx = None

DAY_DATE = {1: '2026-11-05', 2: '2026-11-06', 3: '2026-11-07'}
DATE_DAY = {'5 Nov': 1, '6 Nov': 2, '7 Nov': 3}
DAY_LABEL = {1: '5 Nov', 2: '6 Nov', 3: '7 Nov'}

RE_DAY = re.compile(r'MAMM\s*\d{4}\s*-\s*DAY\s*(\d)', re.I)
RE_SESSION = re.compile(r'(Oral|Poster)\s+Session\s*\((\d+)\)\s*[-\u2013]?\s*', re.I)
RE_ROOM = re.compile(r'\(\s*Rooms?\s*([0-9][0-9,\s&]*)\)', re.I)
RE_TIME = re.compile(r'(\d{1,2}:\d{2})\s*[-~\u2013]\s*(\d{1,2}:\d{2})?')
RE_IDS = re.compile(r'ID_\d+')
RE_TITLE = re.compile(r'(?=\b(?:Dr|Prof|Mr|Ms|Mrs|Assoc\. Prof|Asst\. Prof)\.\s)')


# ------------------------------------------------------------------ helpers
def norm_time(t):
    t = re.sub(r'\s*[~\u2013\u2014-]\s*', '-', (t or '').strip())
    return re.sub(r'-+$', '-', t)


def squash(s):
    return re.sub(r'\s+', ' ', s or '').strip()


def tc_text(tc):
    """Text of one table cell; line breaks and paragraphs become spaces."""
    parts = []
    for p in tc.iter(qn('w:p')):
        buf = []
        for el in p.iter():
            if el.tag == qn('w:t'):
                buf.append(el.text or '')
            elif el.tag in (qn('w:br'), qn('w:cr'), qn('w:tab')):
                buf.append(' ')
        parts.append(''.join(buf))
    return squash(' '.join(parts))


def tc_lines(tc):
    out = []
    for p in tc.iter(qn('w:p')):
        buf = []
        for el in p.iter():
            if el.tag == qn('w:t'):
                buf.append(el.text or '')
            elif el.tag in (qn('w:br'), qn('w:cr')):
                buf.append('\n')
        for ln in ''.join(buf).split('\n'):
            ln = squash(ln)
            if ln:
                out.append(ln)
    return out


def row_cells(tr):
    """[(grid_col, span, vmerge, tc)] for one <w:tr>; vmerge is None/'restart'/'cont'."""
    col, out = 0, []
    for tc in tr.findall(qn('w:tc')):
        span, vm = 1, None
        pr = tc.tcPr
        if pr is not None:
            gs = pr.find(qn('w:gridSpan'))
            if gs is not None:
                span = int(gs.get(qn('w:val')))
            v = pr.find(qn('w:vMerge'))
            if v is not None:
                vm = v.get(qn('w:val')) or 'cont'
        out.append((col, span, vm, tc))
        col += span
    return out


def room_from(text, fallback):
    m = RE_ROOM.search(text)
    if m:
        nums = re.findall(r'\d', m.group(1))
        if len(nums) == 1:
            return 'Room ' + nums[0]
        if len(nums) > 1:
            return 'Rooms ' + ', '.join(nums[:-1]) + ' & ' + nums[-1]
    if re.search(r'Outside\s+Room\s*1', text, re.I):
        return 'Outside Room 1'
    return fallback


def location_room(text):
    t = squash(text)
    if not t or t.lower() == 'location':
        return 'All'
    if re.search(r'Room\s*1\s*,\s*2\s*&\s*3', t):
        return 'All rooms'
    m = re.match(r'^Room\s*(\d)$', t)
    if m:
        return 'Room ' + m.group(1)
    return t


def people(text):
    """'Dr. A Dr. B & Mr. C' -> ['Dr. A', 'Dr. B', 'Mr. C']"""
    text = squash(re.sub(r'^\s*(Session\s+)?Chairs?\s*:?', '', text, flags=re.I))
    text = text.replace(' & ', ' ').replace(' and ', ' ')
    out = [squash(x).strip(' ,;/') for x in RE_TITLE.split(text)]
    return [x for x in out if x]


def join_people(ps):
    if len(ps) <= 1:
        return ps[0] if ps else None
    return ', '.join(ps[:-1]) + ' & ' + ps[-1]


def tidy(text):
    """One schedule cell -> one or more display lines (same rules as the PDF reader)."""
    t = squash(text).strip(' -\u2013')
    t = re.sub(r'(Invited Speech(?: \(\d+\))?)\s*[\u2013-]?\s*(?=(?:Dr\.|Prof\.))', '\\1 \u2013 ', t)
    t = re.sub(r'\s*[\u2013-]\s*\(', ' (', t)
    t = re.sub(r'Session for\s+ISET', 'Session for ISET', t)
    m = re.match(r'^(Invited Speech.*?)\s+(Session for ISET.*)$', t)
    if m:
        return [m.group(1).strip(' -\u2013'), m.group(2).strip(' -\u2013')]
    return [t] if t else []


def name_key(s):
    s = re.sub(r'^(Prof|Dr|Mr|Ms|Mrs)\.\s*(Eng\.\s*)?', '', squash(s), flags=re.I)
    return re.sub(r'[^a-z]', '', s.lower())


# ------------------------------------------------------------- session tables
def parse_session_tables(document):
    sessions = {}
    for tb in document.tables:
        rows = tb._tbl.findall(qn('w:tr'))
        if not rows:
            continue
        first = row_cells(rows[0])
        lines = tc_lines(first[0][3])
        if not lines:
            continue
        m = RE_SESSION.match(lines[0])
        if not m:
            continue
        ncols = max(len(row_cells(r)) for r in rows)
        if ncols < 3:
            continue                     # the two-column summary lists
        kind = 'oral' if m.group(1).lower() == 'oral' else 'poster'
        num = int(m.group(2))
        sid = ('OS' if kind == 'oral' else 'PS') + str(num)
        rest = lines[0][m.end():]
        head = ' '.join(lines)
        chairs, date, time, room = None, None, None, None
        if kind == 'oral':
            name = squash(rest)
            for ln in lines[1:]:
                if re.match(r'^Chairs?\s*:', ln, re.I):
                    chairs = squash(ln.split(':', 1)[1])
                mm = re.match(r'(\d+ \w+),\s*([\d:\-\u2013~ ]+)\s*\((.*)\)', ln)
                if mm:
                    date, time, room = mm.group(1), norm_time(mm.group(2)), room_from('(' + mm.group(3) + ')', mm.group(3))
        else:
            name = 'Poster Session %d' % num
            mm = re.search(r'\((\d+ \w+)\)', rest)
            date = mm.group(1) if mm else None
            ci = re.search(r'Chairs?\s*:', head, re.I)
            if ci:
                chairs = join_people(people(head[ci.end():]))
        ids = []
        for r in rows[1:]:
            cells = row_cells(r)
            if not cells:
                continue
            pid = tc_text(cells[0][3])
            if not pid.startswith('ID_'):
                continue
            prov = '*' in pid
            ids.append((pid.replace('*', '').strip(), prov))
        day = DATE_DAY.get(date)
        sessions[sid] = {
            'id': sid, 'kind': kind, 'num': num, 'name': name,
            'subtitle': name if kind == 'oral' else '',
            'chairs': chairs, 'day': day, 'date': DAY_DATE.get(day), 'dateLabel': date,
            'time': time, 'room': room,
            'paperIds': [i for i, _ in ids],
            '_provisional': dict(ids),
        }
    return sessions


# ---------------------------------------------------------------- day tables
def parse_day_tables(document, sessions, warn):
    days, cur = [], None
    for tb in document.tables:
        rows = tb._tbl.findall(qn('w:tr'))
        if not rows:
            continue
        if not RE_DAY.search(tc_text(row_cells(rows[0])[0][3])):
            continue
        open_merge = {}                         # grid col -> slot dicts still growing
        for tr in rows:
            cells = row_cells(tr)
            txt0 = tc_text(cells[0][3])
            md = RE_DAY.search(txt0)
            if md and len(cells) == 1:
                n = int(md.group(1))
                cur = {'day': n, 'date': DAY_DATE.get(n), 'title': txt0, 'slots': []}
                days.append(cur)
                open_merge = {}
                continue
            if cur is None or len(cells) < 3:
                continue
            time = norm_time(txt0)
            if not RE_TIME.match(time):
                continue
            t_from = time.split('-')[0]
            t_to = time.split('-')[1] if '-' in time else ''
            content, loc, poster = cells[1:-2], cells[-2], cells[-1]
            loc_room = location_room(tc_text(loc[3]))

            # vertical merges: a 'cont' cell stretches the slot above it
            for g, span, vm, tc in cells:
                if vm == 'cont' and g in open_merge:
                    for sl in open_merge[g]:
                        sl['_rows'] += 1
                        if t_to:
                            sl['time'] = sl['time'].split('-')[0] + '-' + t_to
                elif vm != 'cont':
                    open_merge.pop(g, None)

            n = len(content)
            for idx, (g, span, vm, tc) in enumerate(content):
                if vm == 'cont':
                    continue
                text = tc_text(tc)
                if not text:
                    continue
                default = ('Room %d' % (idx + 1)) if n == 3 else room_from(text, loc_room)
                made = emit(cur, sessions, text, time, default, warn)
                if vm == 'restart':
                    for sl in made:
                        sl['_rows'] = 1
                    open_merge[g] = made

            g, span, vm, tc = poster
            ptxt = tc_text(tc)
            if vm != 'cont' and RE_SESSION.search(ptxt):
                emit_poster(cur, sessions, ptxt, time, warn)

    # an "Invited Speech + Session for ISET" cell spanning two time rows:
    # the speech takes the first row, the ISET session the rest
    for d in days:
        for i, sl in enumerate(d['slots']):
            rows = sl.pop('_rows', 1)
            pair = sl.pop('_pair', None)
            if rows > 1 and pair == 'speech':
                a, b = sl['time'].split('-', 1)
                first_end = sl.pop('_first_end', None)
                if first_end:
                    sl['time'] = a + '-' + first_end
                    nxt = d['slots'][i + 1] if i + 1 < len(d['slots']) else None
                    if nxt is not None and nxt.get('_pair') == 'iset':
                        nxt['time'] = first_end + '-' + b
            sl.pop('_first_end', None)
        for sl in d['slots']:
            sl.pop('_pair', None)
            sl.pop('_rows', None)
        # chronological, then Room 1 / 2 / 3, so parallel tracks line up
        d['slots'].sort(key=lambda sl: (sl['time'].split('-')[0].zfill(5), room_rank(sl['room'])))
    return days


def room_rank(r):
    m = re.match(r'^Room (\d)$', r or '')
    return int(m.group(1)) if m else 9


def emit(day, sessions, text, time, default_room, warn):
    made = []
    ms = RE_SESSION.search(text)
    if ms:
        kind = 'oral' if ms.group(1).lower() == 'oral' else 'poster'
        sid = ('OS' if kind == 'oral' else 'PS') + ms.group(2)
        s = sessions.get(sid)
        room = room_from(text, default_room)
        if s is None:
            warn('Word: %s is in the day table but has no session table' % sid)
            name = squash(RE_IDS.sub('', RE_ROOM.sub('', text[ms.end():]))).strip(' ,-\u2013')
        else:
            name = s['name']
            if not s.get('room'):
                s['room'] = room
            if not s.get('time'):
                s['time'] = time
            if not s.get('day'):
                s['day'], s['date'] = day['day'], day['date']
                s['dateLabel'] = DAY_LABEL.get(day['day'])
            if s.get('time') != time or s.get('room') != room or s.get('day') != day['day']:
                warn('Word: %s is at Day %s %s %s in the day table but Day %s %s %s in its own table'
                     % (sid, day['day'], time, room, s.get('day'), s.get('time'), s.get('room')))
            if not s.get('chairs'):
                ci = re.search(r'Session\s+Chairs?\s*:?(.*?)(?=(?:Oral|Poster)\s+Session)', text, re.I)
                if ci:
                    s['chairs'] = join_people(people(ci.group(1)))
        sl = {'time': time, 'room': room,
              'text': 'Oral Session (%s) - %s (%s)' % (ms.group(2), name, room), 'sessionId': sid}
        day['slots'].append(sl)
        made.append(sl)
        return made
    pieces = tidy(text)
    a, b = (time.split('-') + [''])[:2]
    for k, piece in enumerate(pieces):
        sl = {'time': time, 'room': room_from(piece, default_room), 'text': piece}
        if len(pieces) == 2:
            sl['_pair'] = 'speech' if k == 0 else 'iset'
            if k == 0 and b:
                sl['_first_end'] = b
        day['slots'].append(sl)
        made.append(sl)
    return made


def emit_poster(day, sessions, text, row_time, warn):
    ms = RE_SESSION.search(text)
    sid = 'PS' + ms.group(2)
    s = sessions.get(sid)
    after = text[ms.end():]
    mt = RE_TIME.search(after)
    time = norm_time('%s-%s' % mt.groups()) if (mt and mt.group(2)) else row_time
    room = room_from(after, 'Room 1')
    sub = squash(re.split(r'\(\s*Rooms?|\d{1,2}:\d{2}|Session\s+Chair', after)[0]).strip(' -\u2013')
    ci = re.search(r'Session\s+Chairs?\s*:?(.*?)(?=ID_\d|$)', after, re.I)
    if s is None:
        warn('Word: %s is in the day table but has no session table' % sid)
    else:
        s['time'] = time
        s['room'] = room
        s['subtitle'] = sub
        if not s.get('day'):
            s['day'], s['date'] = day['day'], day['date']
            s['dateLabel'] = DAY_LABEL.get(day['day'])
        if not s.get('chairs') and ci:
            s['chairs'] = join_people(people(ci.group(1)))
    day['slots'].append({'time': time, 'room': room,
                         'text': 'Poster Session (%s) - %s (%s)' % (ms.group(2), sub, room)})


# ------------------------------------------------------------------ speakers
def parse_speakers(document, days):
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    out, kind = [], None
    for child in document.element.body.iterchildren():
        if child.tag == qn('w:p'):
            p = Paragraph(child, document)
            t = p.text.strip().upper()
            if p.style.name.startswith('Heading 1'):
                kind = ('keynote' if t.startswith('KEYNOTE') else
                        'invited' if t.startswith('INVITED') else None)
            continue
        if child.tag != qn('w:tbl') or kind is None:
            continue
        rows = child.findall(qn('w:tr'))
        if len(rows) < 2 or len(row_cells(rows[0])) != 1:
            continue
        head = tc_lines(row_cells(rows[0])[0][3])
        if not head:
            continue
        rec = {'kind': kind, 'name': head[0], 'affiliation': ' '.join(head[1:]),
               'topic': '', 'abstract': ''}
        for r in rows[1:]:
            t = tc_text(row_cells(r)[0][3])
            if re.match(r'^Topic\s*:', t, re.I):
                rec['topic'] = squash(t.split(':', 1)[1])
            elif re.match(r'^Abstract\s*:', t, re.I):
                rec['abstract'] = squash(t.split(':', 1)[1])
        key = name_key(rec['name'])
        for d in days:
            for sl in d['slots']:
                if re.search(r'(Keynote|Invited) Speech', sl['text']) and key and key in name_key(
                        re.sub(r'^.*?(?:Speaker:|\u2013)\s*', '', sl['text'])):
                    mnum = re.search(r'Speech \((\d+)\)', sl['text'])
                    rec.update({'day': d['day'], 'date': d['date'], 'time': sl['time'],
                                'room': sl['room'], 'slot': int(mnum.group(1)) if mnum else None})
                    break
            if 'day' in rec:
                break
        out.append(rec)
    return out


# --------------------------------------------------------------------- entry
def has_schedule(path):
    if docx is None:
        return False
    d = docx.Document(path)
    return any(RE_DAY.search(tc_text(row_cells(tb._tbl.findall(qn('w:tr'))[0])[0][3]))
               for tb in d.tables if tb._tbl.findall(qn('w:tr')))


def parse(path, warn=lambda m: None):
    if docx is None:
        sys.exit('python-docx is required:  pip install python-docx')
    d = docx.Document(path)
    sessions = parse_session_tables(d)
    days = parse_day_tables(d, sessions, warn)
    speakers = parse_speakers(d, days)
    provisional = {}
    for s in sessions.values():
        provisional.update(s.pop('_provisional'))
        if s['kind'] == 'oral':
            s['subtitle'] = s['name']
    out = sorted(sessions.values(), key=lambda s: (s['kind'] != 'oral', s['num']))
    return {'days': days, 'sessions': out, 'speakers': speakers, 'provisional': provisional}


if __name__ == '__main__':
    res = parse(sys.argv[1], lambda m: print('  ! ' + m))
    for dd in res['days']:
        print(dd['title'])
        for sl in dd['slots']:
            print('   %-13s %-14s %s' % (sl['time'], sl['room'], sl['text']))
    for s in res['sessions']:
        print('%-5s D%s %-26s %-10s %s | %s | %s' % (s['id'], s['day'], s['time'], s['room'],
                                                   s['name'], s['chairs'], ', '.join(s['paperIds'])))
    for sp in res['speakers']:
        print('%-8s %-28s D%s %s %s | %s' % (sp['kind'], sp['name'], sp.get('day'), sp.get('time'),
                                            sp.get('room'), sp['topic'][:60]))
